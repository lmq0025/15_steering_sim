"""PHASE J1-R2 批量工况：CSV/XLSX 导入、准稳态能力预测、CSV/XLSX 导出。

**同一个物理模型、同一条能力边界，两种计算方式**（任务书第八节）：

``FULL_DYNAMIC``
    单个工况由 :class:`steering_sim.demo.service.DemoService` 真实积分四个通道的
    控制—液压—机械闭环。
``FAST_BATCH_ESTIMATE``
    批量默认：沿指令构型路径逐步求道路负载，用工程液压模型自己的参数核算
    可用推力，某通道需求首次超过可用推力时就停在那里。**这是有来源的准稳态
    作动能力预测，不是一张理想角表**，也绝不冒充完整时间积分。

界面与 ``scripts/j1r2/run_j1r2_batch.py`` 共用本模块，不会出现两套算法。
"""
from __future__ import annotations

from dataclasses import dataclass
import csv
import math
from pathlib import Path
from typing import Any, Callable, Iterable, Mapping, Sequence

import numpy as np

from steering_sim.ackermann.assessment import (OK, VEHICLE_DYNAMIC_INFEASIBLE,
                                               vehicle_feasibility)
from steering_sim.ackermann.ideal import ideal_ackermann
from steering_sim.engineering.mechanism import WHEEL_ORDER
from steering_sim.hydraulics_engineering import load_parameters
from steering_sim.j0_coupling.road_load import road_load

from .scenario import AXLE_COUNT, DemoScenario

ROOT = Path(__file__).resolve().parents[3]
DEFAULT_ENGINEERING_CONFIG = ROOT / "configs" / "j1" / "engineering_config.json"
DEFAULT_ENGINEERING_DESIGN = ROOT / "configs" / "j1" / "engineering_mechanism.json"
DEFAULT_HYDRAULIC_PARAMETERS = ROOT / "configs" / "j1r2" / "engineering_hydraulics.json"

#: 准稳态路径分级数（±30° 上约 1.25° 分辨率）。
PATH_STEPS = 24
DOF_OF_CHANNEL = {"channel_front": "q_front", "channel_rear_1": "q4",
                  "channel_rear_2": "q5", "channel_rear_3": "q6"}
CHANNEL_OF_DOF = {value: key for key, value in DOF_OF_CHANNEL.items()}

#: 导入文件必须含有的列（其余列忽略）。
REQUIRED_COLUMNS = ("target_outer_angle_deg", "speed_kmh")
OPTIONAL_COLUMNS = ("friction_coefficient", "duration_s", "case_id",
                    *(f"axle_mass_{index + 1}_t" for index in range(AXLE_COUNT)))


@dataclass(frozen=True)
class BatchContext:
    """批量计算需要的全部被控对象与能力边界（只建一次）。"""

    model: Any
    adapter: Any
    parameters: Any
    availability_n: Mapping[str, float]
    half_stroke_mm: Mapping[str, float]
    centers: Mapping[str, tuple[float, float]]
    singular_threshold: float

    @property
    def channel_names(self) -> tuple[str, ...]:
        return tuple(self.parameters.channel_names)


def build_context(adapter=None, parameters_path: Path | str | None = None
                  ) -> BatchContext:
    from steering_sim.engineering.model import EngineeringSteeringModel
    from steering_sim.j0_coupling.config import load_j0_config
    from steering_sim.j0_coupling.geometry import J0GeometryAdapter

    if adapter is None:
        model = EngineeringSteeringModel(
            ROOT / "mechanical_model_input_spec_v1.0.yaml",
            DEFAULT_ENGINEERING_CONFIG, DEFAULT_ENGINEERING_DESIGN)
        adapter = J0GeometryAdapter(
            model, load_j0_config(ROOT / "configs" / "j0" / "j0_config.json"),
            mechanical_mode="ENGINEERING_STEERING_MODEL")
    parameters = load_parameters(parameters_path or DEFAULT_HYDRAULIC_PARAMETERS)
    tank = parameters.tank_pressure_pa
    availability = {
        channel.name: ((parameters.supplies[channel.supply].relief_pressure_pa - tank)
                       * channel.cylinder.area_a_m2)
        for channel in parameters.channels}
    return BatchContext(
        model=adapter.model, adapter=adapter, parameters=parameters,
        availability_n=availability,
        half_stroke_mm={channel.name: channel.cylinder.half_stroke_m * 1000.0
                        for channel in parameters.channels},
        centers={wheel: tuple(float(v) for v in adapter.model.centers[wheel])
                 for wheel in WHEEL_ORDER},
        singular_threshold=float(
            adapter.config.mapping["singular_jacobian_threshold_m_per_rad"]))


def estimate_case(context: BatchContext, scenario: DemoScenario,
                  axle_mass_t: float | None = None) -> dict[str, Any]:
    """一个工况的准稳态估计；任何异常都在这里被吸收成结构化结果。"""
    mass = float(axle_mass_t if axle_mass_t is not None
                 else scenario.axle_mass_t[0])
    row: dict[str, Any] = {
        "目标角_deg": scenario.target_outer_angle_deg,
        "车速_kmh": scenario.speed_kmh,
        "摩擦系数": scenario.friction_coefficient,
        "每桥轴荷_t": mass,
        "计算方式": "FAST_BATCH_ESTIMATE",
    }
    try:
        adapter = context.adapter
        model = context.model
        ideal = ideal_ackermann(scenario.target_outer_angle_deg, context.centers)
        for wheel in WHEEL_ORDER:
            row[f"理想_{wheel}"] = round(float(ideal.angles_deg[wheel]), 5)
        geometry = model.solve_geometry(scenario.target_outer_angle_deg)
        state = model.mechanism_state(geometry.q_rad)
        for wheel in WHEEL_ORDER:
            row[f"机构_{wheel}"] = round(float(state.wheel_angles_deg[wheel]), 5)
        row["几何状态"] = geometry.geometry_class
        q_ref = {dof: float(geometry.q_rad[dof]) for dof in DOF_OF_CHANNEL.values()}
        for name in context.channel_names:
            dof = DOF_OF_CHANNEL[name]
            row[f"目标行程_{name}_mm"] = round(
                float(adapter.tables[dof].stroke(q_ref[dof])) * 1000.0, 4)

        reach = {name: 1.0 for name in context.channel_names}
        stopped = {name: False for name in context.channel_names}
        last_force = {name: 0.0 for name in context.channel_names}
        for level in range(1, PATH_STEPS + 1):
            scale = level / PATH_STEPS
            probe = {dof: scale * q_ref[dof] for dof in q_ref}
            road = road_load(adapter, probe, speed_kmh=scenario.speed_kmh,
                             friction_coefficient=scenario.friction_coefficient,
                             axle_mass_t=(mass,) * AXLE_COUNT,
                             curvature_1pm=float(geometry.curvature_1pm),
                             singular_threshold_m_per_rad=context.singular_threshold,
                             use_actual_curvature=True)
            for name in context.channel_names:
                force = float(road.hydraulic_load_force_n[DOF_OF_CHANNEL[name]])
                if not math.isfinite(force):
                    continue
                last_force[name] = force
                if abs(force) > context.availability_n[name] and not stopped[name]:
                    stopped[name] = True
                    reach[name] = (level - 1) / PATH_STEPS
        predicted_q = {dof: reach[name] * q_ref[dof]
                       for name, dof in DOF_OF_CHANNEL.items()}
        predicted = adapter.wheel_angles_deg(predicted_q)
        for wheel in WHEEL_ORDER:
            row[f"液压预测_{wheel}"] = round(float(predicted[wheel]), 5)
        errors = []
        for name in context.channel_names:
            dof = DOF_OF_CHANNEL[name]
            target = abs(float(adapter.tables[dof].stroke(q_ref[dof]))) * 1000.0
            row[f"预测行程_{name}_mm"] = round(abs(float(
                adapter.tables[dof].stroke(predicted_q[dof]))) * 1000.0, 4)
            row[f"行程余量_{name}_mm"] = round(
                context.half_stroke_mm[name] - target, 4)
            row[f"可用推力_{name}_N"] = round(context.availability_n[name], 1)
            row[f"需求推力_{name}_N"] = round(abs(last_force[name]), 1)
            row[f"力裕度_{name}"] = (round(context.availability_n[name]
                                          / abs(last_force[name]), 3)
                                     if abs(last_force[name]) > 1.0 else None)
            row[f"作动能力_{name}"] = "OK" if not stopped[name] else "FORCE_LIMITED"
            row[f"行程状态_{name}"] = ("OK" if target <= context.half_stroke_mm[name]
                                      else "STROKE_OVER")
        for wheel in WHEEL_ORDER:
            errors.append(abs(float(predicted[wheel])
                              - float(state.wheel_angles_deg[wheel])))
        row["预测跟踪误差_deg"] = round(max(errors), 5)

        radius = None
        if ideal.y_icr_m is not None:
            radius = math.hypot(0.0 - ideal.x_icr_m, ideal.y_icr_m)
        feasibility = vehicle_feasibility(scenario.speed_kmh,
                                          scenario.friction_coefficient, radius)
        total_mass = mass * AXLE_COUNT * 1000.0
        row["转弯半径_m"] = None if radius is None else round(radius, 3)
        row["侧向加速度_mps2"] = (None if feasibility.lateral_acceleration_mps2 is None
                                  else round(feasibility.lateral_acceleration_mps2, 4))
        row["附着比"] = (None if feasibility.adhesion_ratio is None
                         else round(feasibility.adhesion_ratio, 4))
        row["需求侧向力_N"] = (None if feasibility.lateral_acceleration_mps2 is None
                               else round(feasibility.lateral_acceleration_mps2
                                          * total_mass, 1))
        row["可提供侧向力_N"] = round(scenario.friction_coefficient * total_mass
                                      * 9.80665, 1)
        row["车辆状态"] = feasibility.status
        limited = any(stopped.values())
        if feasibility.status != OK:
            row["综合状态"] = VEHICLE_DYNAMIC_INFEASIBLE
        elif geometry.geometry_class != "OK":
            row["综合状态"] = geometry.geometry_class
        elif limited:
            row["综合状态"] = "ACTUATOR_CAPABILITY_LIMITED"
        else:
            row["综合状态"] = OK
    except Exception as exc:                      # 单点异常绝不断整批
        row["错误"] = f"{type(exc).__name__}: {exc}"[:200]
        row["几何状态"] = "SOLVER_EXCEPTION"
        row["车辆状态"] = None
        row["综合状态"] = "SOLVER_EXCEPTION"
        for wheel in WHEEL_ORDER:
            row.setdefault(f"理想_{wheel}", None)
            row.setdefault(f"机构_{wheel}", None)
            row.setdefault(f"液压预测_{wheel}", None)
    return row


def run_batch(context: BatchContext, scenarios: Sequence[DemoScenario], *,
              on_progress: Callable[[float], None] | None = None,
              should_stop: Callable[[], bool] | None = None
              ) -> list[dict[str, Any]]:
    """逐个工况做准稳态估计。``should_stop`` 允许界面取消（可取消的后台批处理）。"""
    rows: list[dict[str, Any]] = []
    total = max(1, len(scenarios))
    for index, scenario in enumerate(scenarios, start=1):
        if should_stop is not None and should_stop():
            break
        row = estimate_case(context, scenario)
        row["序号"] = index
        rows.append(row)
        if on_progress is not None and (index % 25 == 0 or index == total):
            try:
                on_progress(index / total)
            except Exception:
                on_progress = None
    return rows


def columns_of(rows: Iterable[Mapping[str, Any]]) -> list[str]:
    columns: list[str] = []
    for row in rows:
        for key in row:
            if key not in columns:
                columns.append(key)
    return columns


def write_csv(rows: Sequence[Mapping[str, Any]], path: Path | str) -> Path:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    columns = columns_of(rows)
    with target.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns)
        writer.writeheader()
        writer.writerows(rows)
    return target


def write_xlsx(rows: Sequence[Mapping[str, Any]], path: Path | str) -> Path | None:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    try:
        from openpyxl import Workbook
    except Exception:                             # openpyxl 缺失时不阻断 CSV
        return None
    columns = columns_of(rows)
    book = Workbook(write_only=True)
    sheet = book.create_sheet(title="J1R2全工况")
    sheet.append(columns)
    for row in rows:
        sheet.append(["" if row.get(key) is None else row.get(key) for key in columns])
    book.save(target)
    return target


# ------------------------------------------------------------------ 导入
def read_cases(path: Path | str) -> list[DemoScenario]:
    """从 CSV / XLSX 读工况。

    必需列：``target_outer_angle_deg``、``speed_kmh``；
    可选列：``friction_coefficient``、``duration_s``、``case_id``、
    ``axle_mass_1_t`` … ``axle_mass_6_t``。缺列时用工程默认值。

    任何非法数值都在这里**显式抛错**（附行号），不静默跳过 —— 用户必须知道
    哪一行没被算进去。
    """
    source = Path(path)
    if not source.exists():
        raise FileNotFoundError(f"工况文件不存在：{source}")
    if source.suffix.lower() in (".xlsx", ".xlsm"):
        rows = _read_xlsx(source)
    else:
        with source.open(encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.DictReader(handle))
    if not rows:
        raise ValueError(f"工况文件没有数据行：{source}")
    missing = [name for name in REQUIRED_COLUMNS if name not in rows[0]]
    if missing:
        raise ValueError(f"工况文件缺少必需列：{', '.join(missing)}")

    from .scenario import ENGINEERING_DEFAULTS, recommended_duration_s
    scenarios: list[DemoScenario] = []
    for index, row in enumerate(rows, start=2):
        try:
            angle = float(row["target_outer_angle_deg"])
            speed = float(row["speed_kmh"])
            mu = float(row.get("friction_coefficient")
                       or ENGINEERING_DEFAULTS["friction_coefficient"])
            duration = row.get("duration_s")
            duration = (float(duration) if duration not in (None, "")
                        else recommended_duration_s(angle))
            loads = []
            for axle in range(1, AXLE_COUNT + 1):
                value = row.get(f"axle_mass_{axle}_t")
                loads.append(float(value) if value not in (None, "")
                             else ENGINEERING_DEFAULTS["axle_mass_t"])
            scenarios.append(DemoScenario(
                angle, speed, friction_coefficient=mu, axle_mass_t=tuple(loads),
                duration_s=duration,
                case_id=str(row.get("case_id") or f"BATCH-{index - 1:04d}")))
        except (TypeError, ValueError) as exc:
            raise ValueError(f"第 {index} 行工况非法：{exc}") from exc
    return scenarios


def _read_xlsx(path: Path) -> list[dict[str, Any]]:
    from openpyxl import load_workbook
    book = load_workbook(path, read_only=True, data_only=True)
    sheet = book.active
    iterator = sheet.iter_rows(values_only=True)
    try:
        header = [str(cell).strip() if cell is not None else "" for cell in next(iterator)]
    except StopIteration:
        return []
    rows = []
    for values in iterator:
        if values is None or all(value in (None, "") for value in values):
            continue
        rows.append({header[i]: values[i] for i in range(min(len(header), len(values)))})
    book.close()
    return rows


def template_rows(count: int = 5) -> list[dict[str, Any]]:
    """导出导入模板（含表头与若干示例行）。"""
    examples = ((0.0, 0.0), (2.0, 5.0), (-10.0, 5.0), (30.0, 10.0), (30.0, 80.0))
    rows = []
    for index in range(max(count, len(examples))):
        angle, speed = examples[index % len(examples)]
        row: dict[str, Any] = {
            "target_outer_angle_deg": angle, "speed_kmh": speed,
            "friction_coefficient": 0.75, "duration_s": "",
            "case_id": f"BATCH-{index + 1:04d}"}
        for axle in range(1, AXLE_COUNT + 1):
            row[f"axle_mass_{axle}_t"] = 11.5
        rows.append(row)
    return rows


def write_template(path: Path | str) -> Path:
    target = Path(path)
    rows = template_rows()
    if target.suffix.lower() in (".xlsx", ".xlsm"):
        written = write_xlsx(rows, target)
        if written is None:
            raise RuntimeError("缺少 openpyxl，无法写出 XLSX 模板")
        return written
    return write_csv(rows, target)

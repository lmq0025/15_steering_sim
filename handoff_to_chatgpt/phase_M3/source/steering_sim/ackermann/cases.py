"""PHASE M2：甲方全工况表批量计算与 CSV/XLSX 导出。

工况网格（甲方要求）：

* 车速 0~80 km/h，每 5 km/h 一档 → 17 档
* 一桥外侧目标角 −30°~+30°，每 1° 一档 → 61 档
* 摩擦系数 0.70 / 0.75 / 0.80 → 3 档

基础车速×角度组合 17×61 = **1037** 组；含摩擦系数共 **3111** 组。

批量只跑**快速机械几何 + 地面载荷 + 适用性评估**，不对 3111 组强制运行
耗时数秒的电液 PID 动态仿真；三模块动态仿真由用户对典型工况单独执行。

每一行都输出：理想阿克曼角、可解时的真实机构角、工程等效角、受载估计角、
五类几何状态、车辆动力学状态与原因。**任何缺失数据都不允许被静默当成 0°**：
无法求解时该列留空并给出原因。
"""
from __future__ import annotations

from dataclasses import dataclass
import csv
import math
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence

from .assessment import (ACKERMANN_TARGET_MISMATCH, KINEMATIC_SOLVER_FAILURE,
                         NUMERICAL_ENVELOPE_LIMIT, OK, PHYSICAL_LINKAGE_LIMIT,
                         VEHICLE_DYNAMIC_INFEASIBLE, CaseAssessment,
                         VehicleFeasibility, assess_wheel_set, direction_violations,
                         vehicle_feasibility)
from .equivalent import CUSTOMER_TAG, EQUIVALENT_TAG, EquivalentChannel
from .ideal import IdealAckermann, ideal_ackermann
from .linkage import REAR_AXLES

WHEEL_ORDER = tuple(f"axle{axle}_{side}" for axle in range(1, 7)
                    for side in ("left", "right"))
AXLES = tuple(range(1, 7))

DEFAULT_SPEEDS_KMH = tuple(float(value) for value in range(0, 81, 5))
DEFAULT_ANGLES_DEG = tuple(float(value) for value in range(-30, 31, 1))
DEFAULT_FRICTION = (0.70, 0.75, 0.80)
DEFAULT_AXLE_LOAD_T = 11.5

#: 机构组 → 通道键。
CHANNEL_OF_DOF = {"q_front": "channel_front", "q4": "channel_rear_1",
                  "q5": "channel_rear_2", "q6": "channel_rear_3"}
DOF_OF_AXLE = {1: "q_front", 2: "q_front", 3: "q3", 4: "q4", 5: "q5", 6: "q6"}


@dataclass(frozen=True, slots=True)
class CaseSpec:
    index: int
    speed_kmh: float
    target_outer_angle_deg: float
    friction_coefficient: float
    axle_loads_t: tuple[float, ...]


def build_case_grid(speeds: Iterable[float] = DEFAULT_SPEEDS_KMH,
                    angles: Iterable[float] = DEFAULT_ANGLES_DEG,
                    friction: Iterable[float] = DEFAULT_FRICTION,
                    axle_loads_t: Sequence[float] | float = DEFAULT_AXLE_LOAD_T
                    ) -> list[CaseSpec]:
    """生成工况网格。默认 17 × 61 × 3 = 3111 组。"""
    if isinstance(axle_loads_t, (int, float)):
        loads = (float(axle_loads_t),) * 6
    else:
        loads = tuple(float(value) for value in axle_loads_t)
        if len(loads) != 6:
            raise ValueError("轴荷必须是 6 个值")
    rows: list[CaseSpec] = []
    index = 0
    for friction_value in friction:
        for speed in speeds:
            for angle in angles:
                index += 1
                rows.append(CaseSpec(index=index, speed_kmh=float(speed),
                                     target_outer_angle_deg=float(angle),
                                     friction_coefficient=float(friction_value),
                                     axle_loads_t=loads))
    return rows


def grid_size(speeds=DEFAULT_SPEEDS_KMH, angles=DEFAULT_ANGLES_DEG,
              friction=DEFAULT_FRICTION) -> tuple[int, int]:
    base = len(tuple(speeds)) * len(tuple(angles))
    return base, base * len(tuple(friction))


# ------------------------------------------------------------------ 单行求值
def _blank() -> dict[str, float]:
    return {wheel: float("nan") for wheel in WHEEL_ORDER}


def _centers_from_model(model) -> dict[str, tuple[float, float]]:
    return {wheel: tuple(float(v) for v in model.centers[wheel])
            for wheel in WHEEL_ORDER}


def _cg_x(model, loads_t: Sequence[float]) -> float:
    weights = [float(value) for value in loads_t]
    total = sum(weights)
    if total <= 0.0:
        return 0.0
    xs = [model.centers[f"axle{axle}_left"][0] for axle in AXLES]
    return float(sum(w * x for w, x in zip(weights, xs)) / total)


def evaluate_case(spec: CaseSpec, model, equivalent: Mapping[str, EquivalentChannel],
                  *, with_loaded: bool = True) -> dict[str, Any]:
    """评估一个工况；返回一行扁平记录。缺失量一律 None，不填 0。"""
    centers = _centers_from_model(model)
    target = spec.target_outer_angle_deg
    row: dict[str, Any] = {
        "序号": spec.index,
        "车速_kmh": spec.speed_kmh,
        "目标角_deg": target,
        "摩擦系数": spec.friction_coefficient,
        "每桥轴荷_t": spec.axle_loads_t[0],
    }
    # 1) 理想阿克曼角（永远可算）
    ideal = ideal_ackermann(target, centers)
    for wheel in WHEEL_ORDER:
        row[f"理想_{wheel}"] = round(ideal.angles_deg[wheel], 6)

    # 2) 整车动力学适用性 a_y = v²/R，R 取质心到 ICR 的距离
    if ideal.y_icr_m is None:
        radius = None
    else:
        cg = _cg_x(model, spec.axle_loads_t)
        radius = math.hypot(cg - ideal.x_icr_m, ideal.y_icr_m)
    vehicle = vehicle_feasibility(spec.speed_kmh, spec.friction_coefficient, radius)
    row["转弯半径_m"] = None if radius is None else round(radius, 4)
    row["侧向加速度_mps2"] = (None if vehicle.lateral_acceleration_mps2 is None
                              else round(vehicle.lateral_acceleration_mps2, 4))
    row["附着比"] = None if vehicle.adhesion_ratio is None else round(vehicle.adhesion_ratio, 4)

    # 3) 真实机构（优先）
    geometry = None
    try:
        geometry = model.solve_geometry(target)
    except Exception as exc:      # 运动学求解失败
        geometry = None
        solver_error = f"{type(exc).__name__}: {exc}"
    else:
        solver_error = ""
    mechanism = _blank()
    loaded = _blank()
    if geometry is not None:
        for wheel in WHEEL_ORDER:
            mechanism[wheel] = float(geometry.theoretical_angle_deg[wheel])
    for wheel in WHEEL_ORDER:
        row[f"机构_{wheel}"] = (None if math.isnan(mechanism[wheel])
                                else round(mechanism[wheel], 6))

    # 4) 工程等效模式（有界、单调、两侧同角）
    equivalent_angles = _blank()
    for axle in REAR_AXLES:
        channel = equivalent[CHANNEL_OF_DOF[f"q{axle}"]]
        if channel.uses_real_linkage:
            for side in ("left", "right"):
                equivalent_angles[f"axle{axle}_{side}"] = mechanism[f"axle{axle}_{side}"]
            continue
        want = ideal.angles_deg[f"axle{axle}_left"]
        magnitude = min(abs(want), channel.wheel_amplitude_max_deg)
        value = math.copysign(magnitude, want) if abs(want) > 0 else 0.0
        for side in ("left", "right"):
            equivalent_angles[f"axle{axle}_{side}"] = value
    for axle in (1, 2, 3):
        for side in ("left", "right"):
            equivalent_angles[f"axle{axle}_{side}"] = mechanism[f"axle{axle}_{side}"]
    for wheel in WHEEL_ORDER:
        row[f"等效_{wheel}"] = (None if math.isnan(equivalent_angles[wheel])
                                else round(equivalent_angles[wheel], 6))

    # 5) 受载估计角（静态柔度修正，非 J0 液压驱动角）
    if with_loaded and geometry is not None:
        try:
            from mechanics_ackermann_v2.model import CaseInput
            case = CaseInput(speed_kmh=spec.speed_kmh,
                             first_axle_outer_angle_deg=target,
                             friction_coefficient=spec.friction_coefficient,
                             axle_mass_t=tuple(spec.axle_loads_t))
            loaded_result = model.solve_loaded_case(case)
            for wheel in WHEEL_ORDER:
                loaded[wheel] = float(
                    loaded_result.actual_loaded_angle_deg[wheel])
        except Exception:
            loaded = _blank()
    for wheel in WHEEL_ORDER:
        row[f"受载_{wheel}"] = (None if math.isnan(loaded[wheel])
                                else round(loaded[wheel], 6))

    # 6) 状态分类
    flags: list[str] = []
    messages: list[str] = []
    geometry_status = OK
    if geometry is None:
        geometry_status = KINEMATIC_SOLVER_FAILURE
        messages.append(f"求解失败：{solver_error}")
        flags.append("KINEMATIC_SOLVER_FAILURE")
    else:
        geometry_status = geometry.geometry_class
        if geometry_status == PHYSICAL_LINKAGE_LIMIT:
            flags.append("LEFT_RIGHT_OPPOSITE")
            messages.append(
                "五、六桥四杆机构为单向可达，同桥左右轮无法同向"
                "（甲方硬点的几何硬约束，非搜索范围问题）。")
        elif geometry_status == ACKERMANN_TARGET_MISMATCH:
            flags.append("ACKERMANN_TARGET_MISMATCH")
            messages.append("机构可运动，但左右轮角无法匹配理想阿克曼角。")
        elif geometry_status == NUMERICAL_ENVELOPE_LIMIT:
            flags.append("NUMERICAL_ENVELOPE_LIMIT")
            messages.append("机构坐标触及数值搜索包络（搜索范围不足，非物理不可达）。")
        # 五类状态是"为什么不可用"的分类；阿克曼匹配度是独立的一个量，
        # 两者都要给，避免让甲方只看到一个笼统结论。
        worst_rms = 0.0
        worst_axle = 0
        for axle in REAR_AXLES:
            values = [geometry.theoretical_angle_deg[f"axle{axle}_{side}"]
                      - geometry.ideal_ackermann_angle_deg[f"axle{axle}_{side}"]
                      for side in ("left", "right")]
            # 两个量本身已经是度，不能再乘一次 math.degrees。
            rms = math.sqrt(sum(v * v for v in values) / len(values))
            if rms > worst_rms:
                worst_rms, worst_axle = rms, axle
        row["阿克曼匹配最大RMS_deg"] = round(worst_rms, 4)
        row["阿克曼匹配最差桥"] = worst_axle
        if worst_rms > 0.25:
            flags.append("ACKERMANN_TARGET_MISMATCH")
            messages.append(
                f"{worst_axle}桥左右轮对理想阿克曼角的 RMS 偏差 "
                f"{worst_rms:.3f}°（限值 0.25°）。")
        for axle in REAR_AXLES:
            left = geometry.theoretical_angle_deg[f"axle{axle}_left"]
            right = geometry.theoretical_angle_deg[f"axle{axle}_right"]
            percent = ("同向" if math.copysign(1.0, left) == math.copysign(1.0, right)
                       else "反号")
            row[f"{axle}桥方向"] = percent
        if not geometry.outer_target_reached:
            flags.append("OUTER_TARGET_NOT_REACHED")
            messages.append("一桥外侧轮未能达到目标转角。")

    row["几何状态"] = geometry_status
    row["车辆状态"] = vehicle.status
    if geometry_status != OK and vehicle.status != OK:
        row["综合状态"] = f"{geometry_status}+{VEHICLE_DYNAMIC_INFEASIBLE}"
    elif geometry_status != OK:
        row["综合状态"] = geometry_status
    else:
        row["综合状态"] = vehicle.status
    if not vehicle.feasible:
        flags.append(VEHICLE_DYNAMIC_INFEASIBLE)
        messages.append(vehicle.message)
    row["说明"] = "；".join(messages)
    row["标记"] = ";".join(flags)
    return row


# --------------------------------------------------------------- 批量驱动
@dataclass(frozen=True, slots=True)
class BatchSummary:
    total: int
    ideal_ok: int
    mechanism_solved: int
    equivalent_used: int
    vehicle_infeasible: int
    status_counts: Mapping[str, int]

    def to_dict(self) -> dict:
        return {
            "total": self.total, "ideal_ok": self.ideal_ok,
            "mechanism_solved": self.mechanism_solved,
            "equivalent_used": self.equivalent_used,
            "vehicle_infeasible": self.vehicle_infeasible,
            "status_counts": dict(self.status_counts),
        }


def run_batch(model, equivalent: Mapping[str, EquivalentChannel],
              specs: Sequence[CaseSpec], *,
              progress=None, with_loaded: bool = True) -> tuple[list[dict], BatchSummary]:
    report = progress or (lambda _done, _total: None)
    rows: list[dict] = []
    counters: dict[str, int] = {}
    solved = 0
    equivalent_used = 0
    vehicle_bad = 0
    for position, spec in enumerate(specs, start=1):
        row = evaluate_case(spec, model, equivalent, with_loaded=with_loaded)
        rows.append(row)
        status = str(row["综合状态"])
        counters[status] = counters.get(status, 0) + 1
        if row.get("机构_axle1_left") is not None:
            solved += 1
        if any(key.startswith("等效_") and value is not None for key, value in row.items()):
            equivalent_used += 1
        if row["车辆状态"] != OK:
            vehicle_bad += 1
        if position % 200 == 0 or position == len(specs):
            report(position, len(specs))
    summary = BatchSummary(
        total=len(rows), ideal_ok=len(rows), mechanism_solved=solved,
        equivalent_used=equivalent_used, vehicle_infeasible=vehicle_bad,
        status_counts=dict(sorted(counters.items(), key=lambda item: -item[1])))
    return rows, summary


# ------------------------------------------------------------------ 导出
def _columns(rows: Sequence[Mapping[str, Any]]) -> list[str]:
    columns: list[str] = []
    for row in rows:
        for key in row:
            if key not in columns:
                columns.append(key)
    return columns


def write_csv(rows: Sequence[Mapping[str, Any]], path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    columns = _columns(rows)
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle)
        writer.writerow(columns)
        for row in rows:
            writer.writerow(["" if row.get(key) is None else row.get(key)
                             for key in columns])
    return path


def write_xlsx(rows: Sequence[Mapping[str, Any]], path: Path,
               *, sheet: str = "全工况表", limit: int | None = None) -> Path:
    """导出 XLSX。默认写出全部行；``limit`` 可用于抽样。"""
    try:
        from openpyxl import Workbook
    except Exception as exc:            # pragma: no cover
        raise RuntimeError(f"未安装 openpyxl，无法导出 XLSX：{exc}") from exc
    submitted = rows if limit is None else rows[:limit]
    columns = _columns(rows)
    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook(write_only=True)
    sheet_obj = book.create_sheet(title=sheet[:31])
    sheet_obj.append(columns)
    for row in submitted:
        sheet_obj.append(["" if row.get(key) is None else row.get(key)
                          for key in columns])
    book.save(path)
    return path


TEMPLATE_COLUMNS = ("序号", "车速_kmh", "一桥外侧目标角_deg", "摩擦系数",
                    "一桥轴荷_t", "二桥轴荷_t", "三桥轴荷_t", "四桥轴荷_t",
                    "五桥轴荷_t", "六桥轴荷_t")


def write_template(path: Path, *, speeds=DEFAULT_SPEEDS_KMH,
                   angles=DEFAULT_ANGLES_DEG, friction=DEFAULT_FRICTION,
                   axle_load_t: float = DEFAULT_AXLE_LOAD_T,
                   limit: int | None = 200) -> Path:
    """生成甲方可填写的工况输入模板（CSV + 同名 XLSX）。"""
    specs = build_case_grid(speeds, angles, friction, axle_load_t)
    if limit is not None:
        specs = specs[:limit]
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle)
        writer.writerow(TEMPLATE_COLUMNS)
        for spec in specs:
            writer.writerow([spec.index, spec.speed_kmh,
                             spec.target_outer_angle_deg, spec.friction_coefficient,
                             *spec.axle_loads_t])
    try:
        from openpyxl import Workbook
        book = Workbook(write_only=True)
        sheet = book.create_sheet(title="工况输入模板")
        sheet.append(list(TEMPLATE_COLUMNS))
        for spec in specs:
            sheet.append([spec.index, spec.speed_kmh, spec.target_outer_angle_deg,
                          spec.friction_coefficient, *spec.axle_loads_t])
        book.save(path.with_suffix(".xlsx"))
    except Exception:
        pass
    return path

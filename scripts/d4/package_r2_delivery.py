"""Build the incremental R2 delivery after truthful test/GUI acceptance."""
from pathlib import Path
import sys,json,hashlib,subprocess,zipfile,io,csv
import xml.etree.ElementTree as ET
ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'outputs/desktop/r2'
TARGET=ROOT/'handoff_to_chatgpt/phase_D4_R2.zip'

def sha(path):return hashlib.sha256(path.read_bytes()).hexdigest()

def main():
    status=json.loads((OUT/'full_repository_run_status.json').read_text('utf-8'))
    if status.get('status')!='PASS' or not status.get('source_hashes_unchanged'):
        raise RuntimeError('Final frozen-source full repository run has not passed')
    if any(sha(ROOT/name)!=digest for name,digest in status['source_hashes_after'].items()):
        raise RuntimeError('Source changed after the final test run')
    suite=ET.parse(OUT/'full_repository.xml').getroot().find('testsuite')
    if int(suite.get('errors',0))+int(suite.get('failures',0)):
        raise RuntimeError('Full repository test failures must be resolved or separately reviewed before final packaging')
    gui=json.loads((OUT/'gui/gui_acceptance.json').read_text('utf-8'))
    if gui.get('status')!='PASS':raise RuntimeError('Real GUI acceptance is not complete')
    baseline=json.loads((OUT/'baseline.json').read_text('utf-8'))
    audit={'protected_hashes':{name:{'baseline':value,'current':sha(ROOT/name),'unchanged':sha(ROOT/name)==value}
                             for name,value in baseline['sha256'].items()},'legacy_xlsx':{},'legacy_csv':{}}
    if not all(v['unchanged'] for v in audit['protected_hashes'].values()):
        raise RuntimeError('Protected baseline changed')
    for name in ('task_spectrum_template.csv','task_spectrum_30day_example.csv'):
        relative=f'data/degradation/{name}'
        original=subprocess.check_output(['git','show',f'HEAD:{relative}'],cwd=ROOT).decode('utf-8-sig')
        current=(ROOT/relative).read_text('utf-8-sig')
        old_reader=csv.DictReader(io.StringIO(original));old_columns=list(old_reader.fieldnames);old_rows=list(old_reader)
        new_reader=csv.DictReader(io.StringIO(current));new_columns=list(new_reader.fieldnames);new_rows=list(new_reader)
        audit['legacy_csv'][relative]={'added_columns':[key for key in new_columns if key not in old_columns],
            'existing_values_unchanged':old_rows==[{k:v for k,v in row.items() if k in old_columns} for row in new_rows],
            'new_rate_values_empty':all(not row.get('seconds_per_degree') for row in new_rows),
            'interpretation':'全仓模板生成测试追加可选速率空列；空值保留历史速率口径，新六秒任务另存r2目录'}
    for name in ('task_spectrum_template.xlsx','task_spectrum_30day_example.xlsx'):
        relative=f'data/degradation/{name}'
        before=subprocess.check_output(['git','show',f'HEAD:{relative}'],cwd=ROOT)
        with zipfile.ZipFile(io.BytesIO(before)) as a,zipfile.ZipFile(ROOT/relative) as b:
            names=set(a.namelist())|set(b.namelist())
            differences=[n for n in sorted(names) if n not in a.namelist() or n not in b.namelist() or a.read(n)!=b.read(n)]
        from openpyxl import load_workbook
        old_book=load_workbook(io.BytesIO(before),read_only=True)
        new_book=load_workbook(ROOT/relative,read_only=True)
        old_rows=list(old_book.active.values);new_rows=list(new_book.active.values)
        width=len(old_rows[0]);old_book.close();new_book.close()
        audit['legacy_xlsx'][relative]={'different_zip_members':differences,
            'existing_cell_values_unchanged':old_rows==[row[:width] for row in new_rows],
            'new_rate_cells_empty':all(all(value is None for value in row[width:]) for row in new_rows[1:]),
            'interpretation':'核心属性时间戳与工作表新增可选空速率列分别审计；非仅元数据变化；新六秒模板另存r2目录'}
    (OUT/'baseline_final_audit.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding='utf-8')
    subprocess.run([sys.executable,str(ROOT/'scripts/d4/save_r2_progress.py'),'--complete'],cwd=ROOT,check=True,stdout=subprocess.DEVNULL)
    progress=json.loads((OUT/'progress.json').read_text('utf-8'));progress['final_zip_created']=True
    progress['delivery']='phase_D4_R2.zip; see adjacent SHA256 and delivery_integrity.json'
    (OUT/'progress.json').write_text(json.dumps(progress,ensure_ascii=False,indent=2),encoding='utf-8')
    files=set()
    def add(path):
        path=Path(path)
        if path.is_file() and '__pycache__' not in path.parts and path.suffix!='.pyc':files.add(path)
    def tree(path,extensions=None):
        for p in Path(path).rglob('*'):
            if extensions is None or p.suffix in extensions:add(p)
    for folder in ('src/steering_sim/desktop','src/steering_sim/degradation','tests/desktop','scripts/d4','data/degradation/r2'):
        tree(ROOT/folder,{'.py','.svg','.ps1','.mjs','.csv','.xlsx','.json','.png'})
    for relative in ('src/steering_sim/j0_coupling/runner.py','src/steering_sim/j0_coupling/angle_reference.py',
        'src/steering_sim/demo/service.py','scripts/d3/run_phase_d3.py','pyproject.toml','requirements-desktop.txt',
        'docs/desktop/phase_D4_R2_report.md','docs/desktop/phase_D4_R2_A_handoff_record.md',
        'configs/desktop/d4_r2_response_analysis.json','tests/degradation/test_d3_rul.py'):
        add(ROOT/relative)
    add(ROOT/'outputs/desktop/working_task_spectrum.csv')
    for name in ('task_spectrum_template','task_spectrum_30day_example'):
        for suffix in ('.csv','.xlsx'):add(ROOT/f'data/degradation/{name}{suffix}')
    for relative in baseline['sha256']:add(ROOT/relative)
    tree(ROOT/'docs/deployment',{'.md','.xlsx'})
    # Reference accuracy evidence, not all exploratory historical timeseries.
    tree(OUT/'a/reference_step_matrix',{'.json','.csv','.txt'})
    tree(OUT/'a/gui_reference_step',{'.json','.csv','.png'})
    tree(OUT/'b',{'.csv','.json'})
    tree(OUT/'cd',{'.csv','.json'})
    for path in OUT.iterdir():
        if path.suffix in ('.json','.log','.xml','.png') and path.name not in ('delivery_integrity.json','delivery_manifest.json'):add(path)
    for path in (OUT/'gui').iterdir():
        if path.suffix in ('.json','.png','.md','.xlsx','.csv'):add(path)
    # Include only the latest completed actual mission of each sample count.
    latest={}
    for metadata in sorted((OUT/'gui/calculations/degradation').glob('*/metadata.json')):
        info=json.loads(metadata.read_text('utf-8'));latest[info['sample_count']]=metadata.parent
    for count in (0,128,1000):
        if count not in latest:raise RuntimeError(f'Missing real GUI forecast: {count}')
        tree(latest[count],{'.csv','.json','.npz'})
    for name,extensions in [('rul_last_run.json',{'.csv','.json'}),('batch_last_run.json',{'.csv','.json'})]:
        source=json.loads((OUT/'gui'/name).read_text('utf-8')).get('output_dir')
        if source:tree(Path(source),extensions)
    logs=sorted((ROOT/'outputs/desktop/logs').glob('desktop_startup*.log'))
    for path in logs[-2:]:add(path)
    manifest={'delivery':'PHASE D4-R2','incremental_package':True,
              'note':'在现有E:/15_project基础上使用；不是经目标机验证的独立EXE',
              'files':[{'path':p.relative_to(ROOT).as_posix(),'bytes':p.stat().st_size,'sha256':sha(p)} for p in sorted(files)]}
    manifest_path=OUT/'delivery_manifest.json'
    manifest_path.write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8');add(manifest_path)
    TARGET.parent.mkdir(parents=True,exist_ok=True)
    with zipfile.ZipFile(TARGET,'w',compression=zipfile.ZIP_DEFLATED,compresslevel=6) as archive:
        for path in sorted(files):archive.write(path,path.relative_to(ROOT).as_posix())
    with zipfile.ZipFile(TARGET) as archive:
        bad=archive.testzip()
        if bad:raise RuntimeError(f'ZIP CRC error: {bad}')
    digest=sha(TARGET)
    TARGET.with_suffix('.zip.sha256').write_text(f'{digest}  {TARGET.name}\n',encoding='ascii')
    integrity={'path':str(TARGET),'bytes':TARGET.stat().st_size,'file_count':len(files),'sha256':digest,'testzip':'PASS'}
    (OUT/'delivery_integrity.json').write_text(json.dumps(integrity,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(integrity,ensure_ascii=False,indent=2))
if __name__=='__main__':main()

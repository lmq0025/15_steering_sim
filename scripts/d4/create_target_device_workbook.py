"""Create the D4-R1 client target-device questionnaire (artifact fallback)."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[2]
OUTPUT = ROOT / "docs" / "deployment" / "TARGET_DEVICE_INFORMATION.xlsx"

ROWS = [
    (1, "设备用途", "", "正式分析机 / 演示机 / 服务器", "是", "确认是否需要桌面GUI"),
    (2, "操作系统", "", "Windows / Linux / 国产Linux", "是", "填写准确名称"),
    (3, "系统版本与补丁", "", "Windows 11 23H2 / UOS V20", "是", "附系统版本截图更佳"),
    (4, "CPU型号", "", "Intel i5-14400 / 飞腾D2000", "是", "填写完整型号"),
    (5, "CPU架构", "", "x86-64 / ARM64 / LoongArch64", "是", "不得只写64位"),
    (6, "CPU核心数", "", "物理核/逻辑线程", "是", "用于性能估计"),
    (7, "内存容量 GB", "", "16 / 32 / 64", "是", "填写可用容量"),
    (8, "可用磁盘空间 GB", "", "至少预留10 GB", "是", "含结果与日志"),
    (9, "图形桌面", "", "Windows桌面 / X11 / Wayland / 无", "是", "无桌面不能运行Qt GUI"),
    (10, "显示分辨率与缩放", "", "1920x1080，125%", "是", "用于界面适配"),
    (11, "显卡与驱动", "", "集成显卡/独显及驱动版本", "是", "Qt/OpenGL兼容"),
    (12, "是否允许安装Python", "", "是 / 否", "是", "决定环境版或目录版"),
    (13, "是否具有管理员权限", "", "是 / 否", "是", "普通用户安装限制"),
    (14, "软件安装限制", "", "白名单/签名/禁止安装", "是", "填写安全策略"),
    (15, "U盘导入限制", "", "允许 / 需审批 / 禁止", "是", "离线交付介质"),
    (16, "杀毒与终端管控", "", "产品名称及版本", "是", "可能影响DLL和启动"),
    (17, "可写目录", "", "建议提供绝对路径", "是", "日志、结果和项目"),
    (18, "中文字体可用性", "", "微软雅黑 / Noto Sans CJK", "是", "界面中文显示"),
    (19, "网络状态", "", "永久离线 / 临时联网 / 内网", "是", "软件本身不需联网"),
    (20, "目标验收联系人", "", "姓名、部门、联系方式", "是", "用于部署确认"),
    (21, "其他限制或备注", "", "自由填写", "否", "包括保密和介质要求"),
]


def build(path: Path = OUTPUT) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "设备信息"
    sheet.sheet_view.showGridLines = False
    sheet.merge_cells("A2:F2")
    sheet["A2"] = "电液转向预测软件目标设备信息采集表"
    sheet.merge_cells("A3:F3")
    sheet["A3"] = "请由甲方设备管理员填写黄色单元格。未确认前不得承诺对应平台可直接运行。"
    sheet["A4"], sheet["B4"] = "必填项状态", "未填写数量"
    sheet["C4"] = '=IF(COUNTBLANK(C7:C26)=0,"信息已完整","仍有必填项未填写")'
    sheet["D4"] = "=COUNTBLANK(C7:C26)"
    headers = ("编号", "信息项", "填写值", "示例或可选值", "必填", "验证说明")
    for column, value in enumerate(headers, 1):
        sheet.cell(6, column, value)
    for row_index, values in enumerate(ROWS, 7):
        for column, value in enumerate(values, 1):
            sheet.cell(row_index, column, value)
    dark = PatternFill("solid", fgColor="245F91")
    input_fill = PatternFill("solid", fgColor="FFF2CC")
    pale = PatternFill("solid", fgColor="EAF2FB")
    thin = Side(style="thin", color="D8E0E7")
    for cell in sheet[6]:
        cell.fill = dark
        cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet["A2"].font = Font(name="Microsoft YaHei", size=16, bold=True, color="17324D")
    sheet["A3"].font = Font(name="Microsoft YaHei", size=10, italic=True, color="5D6B78")
    sheet["A3"].alignment = Alignment(wrap_text=True, vertical="center")
    for row in sheet.iter_rows(min_row=4, max_row=4, min_col=1, max_col=4):
        for cell in row:
            cell.fill = pale
            cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color="17324D")
    for row in sheet.iter_rows(min_row=7, max_row=27, min_col=1, max_col=6):
        for cell in row:
            cell.font = Font(name="Microsoft YaHei", size=10, color="263746")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = Border(bottom=thin)
        row[2].fill = input_fill
        row[0].alignment = Alignment(horizontal="center", vertical="center")
        row[4].alignment = Alignment(horizontal="center", vertical="center")
    sheet.conditional_formatting.add("C7:C26", FormulaRule(
        formula=['LEN(C7)=0'], fill=PatternFill("solid", fgColor="FCE8E6")))
    for address, choices in {
        "C8": '"Windows,Linux,国产Linux,其他"',
        "C11": '"x86-64,ARM64,LoongArch64,其他"',
        "C18:C19": '"是,否,需确认"',
        "C21": '"允许,需审批,禁止"',
    }.items():
        validation = DataValidation(type="list", formula1=choices, allow_blank=True)
        sheet.add_data_validation(validation)
        validation.add(address)
    widths = {"A": 8, "B": 24, "C": 32, "D": 34, "E": 9, "F": 30}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.row_dimensions[2].height = 28
    sheet.row_dimensions[3].height = 32
    sheet.row_dimensions[6].height = 24
    for row in range(7, 28):
        sheet.row_dimensions[row].height = 30
    sheet.freeze_panes = "A7"
    sheet.auto_filter.ref = "A6:F27"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    check = load_workbook(path, data_only=False)
    assert check.sheetnames == ["设备信息"]
    assert check["设备信息"]["C4"].value.startswith("=IF(")
    assert check["设备信息"].max_row == 27
    return path


if __name__ == "__main__":
    print(build())

"""Render a dependency-light visual QA preview of the deployment workbook."""
from pathlib import Path

import matplotlib.pyplot as plt
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "docs" / "deployment" / "TARGET_DEVICE_INFORMATION.xlsx"
OUTPUT = ROOT / "outputs" / "desktop" / "r1" / "target_device_information_preview.png"

book = load_workbook(SOURCE, data_only=False)
sheet = book["设备信息"]
headers = [sheet.cell(6, column).value for column in range(1, 7)]
rows = [[sheet.cell(row, column).value or "" for column in range(1, 7)]
        for row in range(7, min(sheet.max_row, 18) + 1)]

plt.rcParams["font.sans-serif"] = ["Microsoft YaHei", "SimHei", "DejaVu Sans"]
plt.rcParams["axes.unicode_minus"] = False
fig, axis = plt.subplots(figsize=(18, 8), dpi=140)
axis.axis("off")
axis.set_title(sheet["A2"].value, fontsize=18, fontweight="bold", loc="left", pad=18)
table = axis.table(cellText=rows, colLabels=headers, cellLoc="left", colLoc="center",
                   colWidths=[0.05, 0.18, 0.25, 0.24, 0.07, 0.21], loc="upper left")
table.auto_set_font_size(False)
table.set_fontsize(9)
table.scale(1.0, 1.55)
for (row, column), cell in table.get_celld().items():
    if row == 0:
        cell.set_facecolor("#1f4e78")
        cell.get_text().set_color("white")
        cell.get_text().set_fontweight("bold")
    elif column == 2:
        cell.set_facecolor("#fff2cc")
    elif row % 2 == 0:
        cell.set_facecolor("#eef3f8")
OUTPUT.parent.mkdir(parents=True, exist_ok=True)
fig.savefig(OUTPUT, bbox_inches="tight")
plt.close(fig)
print(OUTPUT)
